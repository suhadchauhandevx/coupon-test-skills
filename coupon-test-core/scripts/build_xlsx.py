#!/usr/bin/env python3
"""Build a coupon test-case workbook (SBD layout, manual columns only) from a JSON file.

Usage:
    .venv/bin/python scripts/build_xlsx.py cases.json output.xlsx

This writes a DOCUMENT of test cases. It does not generate or run automation.
Exit code 1 when validation finds errors (the workbook is still written so it can be inspected).
"""
import json
import re
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TYPES = {"P", "N", "E"}
PRIORITIES = {"Critical", "High", "Medium", "Low"}
STATUSES = ["Not Run", "Pass", "Fail", "Blocked", "Query"]
HEADERS = ["ID", "Title", "Screen", "Route", "Type", "Priority", "Tags",
           "Preconditions", "Steps", "Expected Result", "Status"]
WIDTHS = [15, 46, 10, 22, 7, 10, 22, 36, 56, 56, 10]
TC_SHEET = "Test Cases"
NCOLS = len(HEADERS)

NAVY, BLUE, BANNER, ZEBRA, WARN = "1F2D3D", "2563EB", "1F4E79", "F8FAFC", "FFF4CE"
INK = "1F2D3D"
thin = Side(style="thin", color="D0D5DD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(bold=False, size=9, color=INK):
    return Font(name="Arial", bold=bold, size=size, color=color)


def fill(rgb):
    return PatternFill("solid", fgColor=rgb)


def align(h="left", v="top"):
    return Alignment(horizontal=h, vertical=v, wrap_text=True)


def banner(ws, row, text, ncols, rgb, size, height):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.fill, c.alignment = font(True, size, "FFFFFF"), fill(rgb), align(v="center")
    for col in range(2, ncols + 1):
        ws.cell(row=row, column=col).fill = fill(rgb)
    ws.row_dimensions[row].height = height


def row_height(case):
    lines = max(
        str(case.get("steps", "")).count("\n") + 1,
        len(str(case.get("expected", ""))) // 60 + 1,
        len(str(case.get("pre", ""))) // 38 + 1,
        len(str(case.get("title", ""))) // 48 + 1,
    )
    return max(30, min(11.5 * lines + 6, 400))


def validate(data):
    errs, warns = [], []
    seen = Counter()
    for sec in data["sections"]:
        for c in sec["cases"]:
            seen[c["id"]] += 1
            where = c["id"]
            if c.get("type") not in TYPES:
                errs.append(f"{where}: type must be P/N/E, got {c.get('type')!r}")
            if c.get("priority") not in PRIORITIES:
                errs.append(f"{where}: priority must be one of {sorted(PRIORITIES)}, got {c.get('priority')!r}")
            for k in ("title", "steps", "expected"):
                if not str(c.get(k, "")).strip():
                    errs.append(f"{where}: empty {k}")
            exp = str(c.get("expected", ""))
            if re.search(r"\b(either|or else)\b", exp, re.I) or re.search(r"\.\s*or\b", exp):
                warns.append(f"{where}: Expected Result may allow two outcomes ('either'/'or'); "
                             "pick one and move the decision to Open Items")
            if not re.match(r"^1\.", str(c.get("steps", "")).strip()):
                warns.append(f"{where}: steps should be numbered starting at '1.'")
    for i, n in seen.items():
        if n > 1:
            errs.append(f"{i}: duplicate ID ({n} times)")
    return errs, warns


def build(data, out_path):
    wb = Workbook()
    meta = data["meta"]
    sections = data["sections"]
    total = sum(len(s["cases"]) for s in sections)

    # ---------- Cover & Legend ----------
    ws = wb.active
    ws.title = "Cover & Legend"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 108
    banner(ws, 1, f"{meta['project']} — {meta['coupon_type']} | Manual Test Suite | {meta.get('version', 'v1.0')}", 2, NAVY, 13, 27.75)
    r = 3

    def heading(text):
        nonlocal r
        banner(ws, r, text, 2, BANNER, 11, 21.75)
        r += 1

    def kv(label, text, warn=False):
        nonlocal r
        a = ws.cell(row=r, column=1, value=label)
        a.font, a.alignment = font(True), align(v="center")
        b = ws.cell(row=r, column=2, value=text)
        b.font, b.alignment = font(), align()
        if warn:
            b.fill = fill(WARN)
        ws.row_dimensions[r].height = max(18, 12 * (len(str(text)) // 120 + 1) + 4)
        r += 1

    heading("About this version")
    for label, key in [("Client", "client"), ("Scope", "scope"), ("Not covered here", "not_covered"),
                       ("Built from", "built_from"), ("What changed", "what_changed")]:
        if meta.get(key):
            kv(label, meta[key])
    kv("Total cases", f"{total} across {len(sections)} sections")
    r += 1
    heading("Column guide")
    for label, text in [
        ("ID", "Stable identifier. Quote it in bug reports."),
        ("Screen / Route", "Section code and the page or flow the case runs against."),
        ("Type", "P = Positive, N = Negative, E = Edge case."),
        ("Priority", "Critical / High / Medium / Low — drives smoke vs full-regression selection."),
        ("Tags", "Suite selectors: smoke, sanity, regression, negative, boundary, edge, integrity, rbac, "
                 "gap-coverage (never executed manually before), defect-suspect (expected to fail today)."),
        ("Preconditions", "State the system must be in before step 1."),
        ("Steps", "Numbered manual steps."),
        ("Expected Result", "The single assertion outcome. Where behaviour is a known defect it says so explicitly."),
        ("Status", "Execution result: Not Run / Pass / Fail / Blocked / Query. Query = waiting on an Open Item."),
    ]:
        kv(label, text)
    r += 1
    if meta.get("safety_rule"):
        heading("Test-data safety rule")
        kv("Rule", meta["safety_rule"], warn=True)
        r += 1
    if meta.get("notes"):
        heading("Known behaviours to watch")
        for n in meta["notes"]:
            kv(n["label"], n["text"], warn=n.get("warn", False))

    # ---------- Test Cases ----------
    tc = wb.create_sheet(TC_SHEET)
    banner(tc, 1, f"{meta['project']} — {meta['coupon_type']} | {meta.get('version', 'v1.0')} | {total} test cases | manual",
           NCOLS, NAVY, 13, 27.75)
    for i, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
        c = tc.cell(row=2, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = font(True, 10, "FFFFFF"), fill(BLUE), align(v="center"), BORDER
        tc.column_dimensions[get_column_letter(i)].width = w
    tc.row_dimensions[2].height = 24
    row = 3
    banner_rows = []
    for sec in sections:
        n = len(sec["cases"])
        banner(tc, row, f"{sec['name'].upper()}  ({sec['code']})   —   {sec['route']}   —   {n} cases",
               NCOLS, BANNER, 11, 21.75)
        banner_rows.append(row)
        row += 1
        for k, c in enumerate(sec["cases"]):
            vals = [c["id"], c["title"], sec["code"], c.get("route", sec["route"]), c["type"], c["priority"],
                    c.get("tags", "regression"), c.get("pre", "—"), c["steps"], c["expected"],
                    c.get("status", "Not Run")]
            for col, v in enumerate(vals, start=1):
                cell = tc.cell(row=row, column=col, value=v)
                cell.border = BORDER
                centered = col in (3, 5, 6, 11)
                cell.alignment = align("center" if centered else "left", "center" if centered else "top")
                cell.font = font(bold=(col in (1, 2)))
                if k % 2 == 1:
                    cell.fill = fill(ZEBRA)
            tc.row_dimensions[row].height = row_height(c)
            row += 1
    last = row - 1
    tc.freeze_panes = "B3"
    tc.auto_filter.ref = f"A2:{get_column_letter(NCOLS)}{last}"
    for col, vals in (("E", '"P,N,E"'), ("F", '"Critical,High,Medium,Low"'),
                      ("K", '"' + ",".join(STATUSES) + '"')):
        dv = DataValidation(type="list", formula1=vals, allow_blank=True)
        tc.add_data_validation(dv)
        dv.add(f"{col}3:{col}{last}")

    # ---------- Summary (live formulas) ----------
    sm = wb.create_sheet("Summary")
    heads = ["Module", "Screen", "Route", "P", "N", "E", "Total", "Critical"]
    for i, w in enumerate([46, 10, 26, 7, 7, 7, 8, 9], start=1):
        sm.column_dimensions[get_column_letter(i)].width = w
    banner(sm, 1, f"{meta['coupon_type']} — Summary by section (live formulas over '{TC_SHEET}')", 8, NAVY, 13, 27.75)
    for i, h in enumerate(heads, start=1):
        c = sm.cell(row=2, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = font(True, 10, "FFFFFF"), fill(BLUE), align("center", "center"), BORDER
    rng = lambda col: f"'{TC_SHEET}'!${col}$3:${col}${last}"
    sr = 3
    for sec in sections:
        sm.cell(row=sr, column=1, value=sec["name"])
        sm.cell(row=sr, column=2, value=sec["code"])
        sm.cell(row=sr, column=3, value=sec["route"])
        for col, t in ((4, "P"), (5, "N"), (6, "E")):
            sm.cell(row=sr, column=col, value=f'=COUNTIFS({rng("C")},"{sec["code"]}",{rng("E")},"{t}")')
        sm.cell(row=sr, column=7, value=f"=SUM(D{sr}:F{sr})")
        sm.cell(row=sr, column=8, value=f'=COUNTIFS({rng("C")},"{sec["code"]}",{rng("F")},"Critical")')
        for col in range(1, 9):
            c = sm.cell(row=sr, column=col)
            c.font, c.border = font(bold=(col == 1)), BORDER
            c.alignment = align("left" if col <= 3 else "center", "center")
        sr += 1
    first, lastsec = 3, sr - 1
    sm.cell(row=sr, column=1, value="TOTAL")
    for col in range(4, 9):
        L = get_column_letter(col)
        sm.cell(row=sr, column=col, value=f"=SUM({L}{first}:{L}{lastsec})")
    for col in range(1, 9):
        c = sm.cell(row=sr, column=col)
        c.font, c.fill, c.border = font(True), fill(ZEBRA), BORDER
        c.alignment = align("left" if col <= 3 else "center", "center")
    sr += 2
    banner(sm, sr, "Execution status", 8, BANNER, 11, 21.75)
    sr += 1
    for s in STATUSES:
        sm.cell(row=sr, column=1, value=s).font = font(True)
        c = sm.cell(row=sr, column=4, value=f'=COUNTIF({rng("K")},"{s}")')
        c.font, c.alignment = font(), align("center", "center")
        sr += 1
    sm.cell(row=sr + 1, column=1, value="P = Positive  ·  N = Negative  ·  E = Edge case").font = font()

    # ---------- Open items ----------
    oi = wb.create_sheet("Mismatches & Open Items")
    for i, w in enumerate([5, 22, 44, 52, 56], start=1):
        oi.column_dimensions[get_column_letter(i)].width = w
    banner(oi, 1, "Open items & assumptions — need a decision from the team before the dependent cases can pass/fail",
           5, NAVY, 13, 27.75)
    for i, h in enumerate(["#", "Area", "Documented / expected", "Observed / unknown", "Impact & what I need from the team"], 1):
        c = oi.cell(row=2, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = font(True, 10, "FFFFFF"), fill(BLUE), align(v="center"), BORDER
    oi.freeze_panes = "A3"
    for n, it in enumerate(data.get("open_items", []), start=1):
        for col, v in enumerate([n, it["area"], it["documented"], it["observed"], it["impact"]], start=1):
            c = oi.cell(row=2 + n, column=col, value=v)
            c.font, c.border, c.alignment = font(bold=(col == 2)), BORDER, align()
        oi.row_dimensions[2 + n].height = max(30, 12 * (max(len(it["observed"]), len(it["impact"])) // 55 + 1) + 6)

    wb.save(out_path)
    return total, last


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    data = json.load(open(sys.argv[1], encoding="utf-8"))
    errs, warns = validate(data)
    total, last = build(data, sys.argv[2])
    c = Counter(x["type"] for s in data["sections"] for x in s["cases"])
    p = Counter(x["priority"] for s in data["sections"] for x in s["cases"])
    print(f"wrote {sys.argv[2]}: {total} cases | types {dict(c)} | priorities {dict(p)}")
    for w in warns:
        print("WARN ", w)
    for e in errs:
        print("ERROR", e)
    sys.exit(1 if errs else 0)


if __name__ == "__main__":
    main()
